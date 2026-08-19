#!/usr/bin/env python3
"""Build dark-theme Samsung Science CTC-EMT PPTX and slide schematics."""

from __future__ import annotations

from pathlib import Path
import math

import matplotlib.pyplot as plt
from matplotlib import font_manager as fm
from matplotlib.patches import Circle, FancyArrowPatch, FancyBboxPatch, Rectangle
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "outputs" / "figures"
DECK = ROOT / "outputs" / "deck"

BG = "#090d14"
PANEL = "#111827"
PANEL2 = "#172033"
TEXT = "#e5e7eb"
MUTED = "#9ca3af"
CYAN = "#22d3ee"
PURPLE = "#a78bfa"
ORANGE = "#fb923c"
GREEN = "#34d399"
RED = "#fb7185"
GOLD = "#fbbf24"

KR_FONT = Path("/home/seungho/.local/share/fonts/NanumGothic-Regular.ttf")
KR_FONT_BOLD = Path("/home/seungho/.local/share/fonts/NanumGothic-Bold.ttf")
if KR_FONT.exists():
    fm.fontManager.addfont(str(KR_FONT))
    if KR_FONT_BOLD.exists():
        fm.fontManager.addfont(str(KR_FONT_BOLD))
    plt.rcParams["font.family"] = "NanumGothic"


def setup_ax(path: Path, title: str | None = None):
    fig, ax = plt.subplots(figsize=(16, 9), dpi=180)
    fig.patch.set_facecolor(BG)
    ax.set_facecolor(BG)
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 9)
    ax.axis("off")
    if title:
        ax.text(0.7, 8.25, title, color=TEXT, fontsize=24, fontweight="bold")
    return fig, ax


def save(fig, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, facecolor=BG, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)


def box(ax, x, y, w, h, label, color=CYAN, sub=None, fs=16):
    rect = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.03,rounding_size=0.12",
        linewidth=1.6,
        edgecolor=color,
        facecolor=PANEL,
        alpha=0.96,
    )
    ax.add_patch(rect)
    ax.text(x + w / 2, y + h * 0.58, label, color=TEXT, fontsize=fs, ha="center", va="center", fontweight="bold")
    if sub:
        ax.text(x + w / 2, y + h * 0.25, sub, color=MUTED, fontsize=fs - 4, ha="center", va="center")


def arrow(ax, x1, y1, x2, y2, color=CYAN, lw=2.4):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=18, color=color, lw=lw))


def slide01():
    fig, ax = setup_ax(FIG / "slide01_problem.png")
    stages = [
        ("Diagnosis", "US/FNA"),
        ("Thyroidectomy", "source removed"),
        ("Surveillance", "Tg/US"),
        ("Uncertainty", "residual-risk state"),
    ]
    xs = [1.3, 5.1, 8.9, 12.7]
    for i, (t, s) in enumerate(stages):
        box(ax, xs[i], 4.3, 2.4, 1.25, t, [CYAN, PURPLE, GREEN, ORANGE][i], s)
        if i < 3:
            arrow(ax, xs[i] + 2.45, 4.95, xs[i + 1] - 0.2, 4.95, MUTED)
    ax.text(0.8, 7.5, "PTC is curable, but residual-risk biology is invisible after surgery", color=TEXT, fontsize=25, fontweight="bold")
    ax.text(0.8, 2.55, "Current tools are indirect or delayed: pathology, Tg/anti-Tg, ultrasound, recurrence surveillance.", color=MUTED, fontsize=16)
    ax.text(0.8, 1.9, "Missing readout: serial cellular state + matched genetic identity in blood.", color=CYAN, fontsize=18, fontweight="bold")
    save(fig, FIG / "slide01_problem.png")


def slide02():
    fig, ax = setup_ax(FIG / "slide02_hypothesis.png")
    ax.text(0.8, 7.5, "Thyroidectomy as a human perturbation model", color=TEXT, fontsize=26, fontweight="bold")
    ax.text(0.8, 6.95, "Surgery should collapse shedding; persistence or EMT-state shift after surgery is the signal.", color=MUTED, fontsize=16)
    for i, (name, color, y) in enumerate([("E", CYAN, 5.3), ("E/M", PURPLE, 4.2), ("M", ORANGE, 3.1)]):
        ax.add_patch(Circle((2.1, y), 0.42, color=color, alpha=0.9))
        ax.text(2.1, y, name, ha="center", va="center", color=BG, fontsize=16, fontweight="bold")
        ax.add_patch(Circle((12.3, y - 0.2 + i * 0.1), 0.30, color=color, alpha=0.45 if i == 0 else 0.95))
        ax.text(12.3, y - 0.2 + i * 0.1, name, ha="center", va="center", color=BG, fontsize=12, fontweight="bold")
    box(ax, 5.7, 3.65, 3.5, 1.5, "Thyroidectomy", PURPLE, "controlled perturbation", fs=18)
    arrow(ax, 2.8, 4.2, 5.5, 4.4, MUTED)
    arrow(ax, 9.25, 4.4, 11.75, 4.3, MUTED)
    ax.text(1.25, 6.05, "Pre-op baseline", color=TEXT, fontsize=16, fontweight="bold")
    ax.text(11.0, 6.05, "Post-op signal", color=TEXT, fontsize=16, fontweight="bold")
    ax.text(10.25, 2.05, "Persistent EM/M + tumor-matched genotype", color=ORANGE, fontsize=18, fontweight="bold")
    ax.text(10.25, 1.55, "Residual-risk biology hypothesis", color=MUTED, fontsize=14)
    save(fig, FIG / "slide02_hypothesis.png")


def slide03():
    fig, ax = setup_ax(FIG / "slide03_literature_gap.png")
    ax.text(0.7, 7.6, "CTC and EMT phenotypes are measurable in PTC", color=TEXT, fontsize=26, fontweight="bold")
    headers = ["Study", "What exists", "Gap"]
    xs = [0.8, 5.1, 10.6]
    for x, h in zip(xs, headers):
        ax.text(x, 6.75, h, color=CYAN, fontsize=16, fontweight="bold")
    rows = [
        ("Yu 2024", "62 PTC; pre-op, 2w, 3m;\nCTC 87%; EMT/M predominant;\npost-op count declines", "No matched CTC/cfDNA\nNGS clone map"),
        ("Li 2022", "394 thyroid cancer;\nCTC EMT subtyping;\nM-CTC prognostic signal", "Mixed histologies;\nnot serial thyroidectomy NGS"),
        ("Sato 2021", "PTC BRAF ctDNA before/after\nsurgery by ddPCR", "Single mutation;\nsmall N; not CTC state"),
    ]
    y = 5.8
    for i, row in enumerate(rows):
        ax.add_patch(Rectangle((0.55, y - 0.85), 14.9, 1.18, facecolor=PANEL if i % 2 == 0 else PANEL2, edgecolor="#263244"))
        for x, txt in zip(xs, row):
            ax.text(x, y, txt, color=TEXT if x == xs[0] else MUTED, fontsize=13, va="center", fontweight="bold" if x == xs[0] else None)
        y -= 1.45
    ax.text(0.8, 1.25, "Proposal gap: serial PTC CTC-EMT phenotype + matched tumor-normal NGS + postoperative blood genetic map.", color=ORANGE, fontsize=17, fontweight="bold")
    save(fig, FIG / "slide03_literature_gap.png")


def slide04():
    fig, ax = setup_ax(FIG / "slide04_missing_dataset.png")
    ax.text(0.8, 7.55, "No public dataset has the full serial CTC-EMT-NGS map", color=TEXT, fontsize=25, fontweight="bold")
    items = [
        "serial pre/post thyroidectomy blood",
        "CTC EMT phenotype transition",
        "matched tumor-normal NGS",
        "cfDNA or CTC-enriched NGS",
        "postoperative outcome/follow-up",
    ]
    for i, item in enumerate(items):
        y = 6.35 - i * 0.9
        ax.add_patch(Rectangle((1.1, y - 0.28), 0.42, 0.42, edgecolor=RED, facecolor="none", lw=2))
        ax.plot([1.12, 1.5], [y - 0.25, y + 0.13], color=RED, lw=2)
        ax.plot([1.5, 1.12], [y - 0.25, y + 0.13], color=RED, lw=2)
        ax.text(1.85, y, item, color=TEXT, fontsize=18, va="center")
    box(ax, 10.0, 2.5, 4.5, 2.2, "New dataset", ORANGE, "created by this proposal", fs=20)
    ax.text(10.45, 1.55, "This absence is the justification,\nnot a weakness to hide.", color=GOLD, fontsize=16, fontweight="bold")
    save(fig, FIG / "slide04_missing_dataset.png")


def slide05():
    fig, ax = setup_ax(FIG / "slide05_driver_not_enough.png")
    ax.text(0.7, 7.55, "Driver mutation does not capture the full residual-risk state", color=TEXT, fontsize=24, fontweight="bold")
    labels = ["Largest BRAF label", "BRAF ambiguity"]
    vals = [33.2, 66.8]
    ax.barh([4.8, 3.6], vals, color=[CYAN, ORANGE], height=0.55)
    for y, val in zip([4.8, 3.6], vals):
        ax.text(val + 2, y, f"{val:.1f}%", color=TEXT, fontsize=18, va="center", fontweight="bold")
    ax.set_xlim(0, 100)
    ax.text(0, 5.65, "TCGA-THCA BRAF-mutant tumors: 274 patients, 7 vulnerability labels", color=MUTED, fontsize=16)
    ax.text(0, 2.45, "Driver eta-squared: RAI 30.9%, HLA-I/APM 15.6%, immune visibility 12.5%, CD8 exclusion 16.2%", color=MUTED, fontsize=14)
    ax.text(0, 1.45, "Boundary: supports marker logic and genotype insufficiency, not CTC biology.", color=GOLD, fontsize=15, fontweight="bold")
    save(fig, FIG / "slide05_driver_not_enough.png")


def slide06():
    fig, ax = setup_ax(FIG / "slide06_spatial_context.png")
    ax.set_ylim(0, 60)
    ax.text(
        0.045,
        0.90,
        "Spatial tissue states nominate where EMT/CTC signals may originate",
        color=TEXT,
        fontsize=22,
        fontweight="bold",
        transform=ax.transAxes,
    )
    ax.text(0.06, 0.78, "GSE250521: 16 slides, 57,144 spots", color=MUTED, fontsize=16, transform=ax.transAxes)
    ax.text(
        0.06,
        0.72,
        "16/16 slides: same-niche coherence z > 2",
        color=CYAN,
        fontsize=18,
        fontweight="bold",
        transform=ax.transAxes,
    )
    xs = [3, 6.5, 10]
    vals = [6.09, 22.00, 43.16]
    labs = ["min z", "median z", "max z"]
    ax.bar(xs, vals, width=1.4, color=[CYAN, PURPLE, ORANGE])
    for x, v, lab in zip(xs, vals, labs):
        ax.text(x, v + 1.5, f"{v:.2f}", color=TEXT, fontsize=18, ha="center", fontweight="bold")
        ax.text(x, 2.4, lab, color=MUTED, fontsize=14, ha="center")
    save(fig, FIG / "slide06_spatial_context.png")


def slide07():
    fig, ax = setup_ax(FIG / "slide07_ctc_ngs_panel.png")
    ax.text(0.7, 7.55, "From tissue state to circulating state", color=TEXT, fontsize=26, fontweight="bold")
    modules = [
        ("Epithelial", "EpCAM, KRT8/18/19, CDH1", CYAN),
        ("Mesenchymal", "VIM, FN1, CDH2, SNAI2, ZEB1, TWIST1", ORANGE),
        ("Thyroid identity", "PAX8, TG, TPO, TSHR, SLC5A5", PURPLE),
        ("Survival/stemness", "CD44, ALDH1A1, SOX9, AXL, MET", GREEN),
        ("Immune/stress", "CD274, CD47, HLA/B2M/TAP, CA9, VEGFA", GOLD),
    ]
    y = 6.25
    for name, genes, color in modules:
        box(ax, 1.0, y - 0.35, 4.1, 0.7, name, color, fs=14)
        ax.text(5.5, y, genes, color=TEXT, fontsize=14, va="center")
        y -= 1.0
    ax.text(1.0, 1.1, "Panel purpose: define E, E/M, M CTC state and anchor blood signal to tumor genotype.", color=CYAN, fontsize=16, fontweight="bold")
    save(fig, FIG / "slide07_ctc_ngs_panel.png")


def slide08():
    fig, ax = setup_ax(FIG / "slide08_sampling_timeline.png")
    ax.text(0.75, 7.6, "Serial blood around thyroidectomy", color=TEXT, fontsize=26, fontweight="bold")
    times = [("T0", "pre-op"), ("T1", "2 weeks"), ("T2", "3 months"), ("T3", "6-12 months\nor trigger")]
    xs = [2.0, 5.8, 9.6, 13.4]
    ax.plot([xs[0], xs[-1]], [4.7, 4.7], color=MUTED, lw=3)
    for i, ((t, desc), x) in enumerate(zip(times, xs)):
        ax.add_patch(Circle((x, 4.7), 0.45, color=[CYAN, PURPLE, GREEN, ORANGE][i]))
        ax.text(x, 4.7, t, color=BG, fontsize=16, fontweight="bold", ha="center", va="center")
        ax.text(x, 3.85, desc, color=TEXT, fontsize=14, ha="center")
    box(ax, 5.1, 5.65, 3.3, 1.0, "Thyroidectomy", PURPLE, "perturbation", fs=16)
    ax.text(1.0, 1.6, "Per timepoint: CTC phenotype + cfDNA. Tissue: matched tumor-normal NGS + FFPE/mIHC.", color=CYAN, fontsize=16, fontweight="bold")
    save(fig, FIG / "slide08_sampling_timeline.png")


def slide09():
    fig, ax = setup_ax(FIG / "slide09_genetic_map.png")
    ax.text(0.7, 7.55, "Matched tissue NGS anchors the liquid-biopsy signal", color=TEXT, fontsize=25, fontweight="bold")
    box(ax, 1.0, 4.8, 3.0, 1.2, "Tumor tissue", PURPLE, "matched normal NGS")
    box(ax, 6.0, 4.8, 3.5, 1.2, "Genetic map", CYAN, "driver/private variants")
    box(ax, 12.0, 5.6, 2.8, 1.0, "cfDNA", GREEN, "serial persistence")
    box(ax, 12.0, 3.8, 2.8, 1.0, "CTC-enriched", ORANGE, "QC subset")
    arrow(ax, 4.1, 5.4, 5.8, 5.4, MUTED)
    arrow(ax, 9.6, 5.4, 11.75, 6.05, MUTED)
    arrow(ax, 9.6, 5.15, 11.75, 4.3, MUTED)
    ax.text(1.0, 2.05, "CTC phenotype alone can be nonspecific. NGS asks: does postoperative blood match the tumor clone?", color=GOLD, fontsize=16, fontweight="bold")
    save(fig, FIG / "slide09_genetic_map.png")


def slide10():
    fig, ax = setup_ax(FIG / "slide10_claim_boundary.png")
    ax.text(0.7, 7.55, "What we will test, and what we will not claim", color=TEXT, fontsize=25, fontweight="bold")
    left = ["CTC E/EM/M transition", "tissue genotype concordance", "cfDNA persistence", "association with LVI/LNM/ETE/Tg/US"]
    right = ["ready clinical diagnostic", "recurrence prediction before validation", "vendor technology as novelty", "CTC equals viable metastasis"]
    ax.text(1.0, 6.45, "Allowed tests", color=GREEN, fontsize=18, fontweight="bold")
    ax.text(8.8, 6.45, "Forbidden claims", color=RED, fontsize=18, fontweight="bold")
    for i, t in enumerate(left):
        ax.text(1.1, 5.75 - i * 0.75, f"+ {t}", color=TEXT, fontsize=15)
    for i, t in enumerate(right):
        ax.text(8.9, 5.75 - i * 0.75, f"- {t}", color=TEXT, fontsize=15)
    ax.add_patch(Rectangle((0.7, 2.2), 6.7, 4.65, fill=False, edgecolor=GREEN, lw=1.6))
    ax.add_patch(Rectangle((8.4, 2.2), 6.7, 4.65, fill=False, edgecolor=RED, lw=1.6))
    ax.text(1.0, 1.0, "Primary endpoint: biological state transition, not clinical deployment.", color=CYAN, fontsize=17, fontweight="bold")
    save(fig, FIG / "slide10_claim_boundary.png")


def slide11():
    fig, ax = setup_ax(FIG / "slide11_budget_roadmap.png")
    ax.text(0.7, 7.55, "Stage-gated execution protects scientific rigor and budget", color=TEXT, fontsize=24, fontweight="bold")
    years = [("Year 1", "assay lock\nIRB/cohort\npilot NGS"), ("Year 2", "scale cohort\ncfDNA/CTC feasibility\nmIHC validation"), ("Year 3", "transition model\ngenetic map\nmanuscript/IP")]
    xs = [1.2, 6.0, 10.8]
    for i, (yr, desc) in enumerate(years):
        box(ax, xs[i], 4.85, 3.6, 1.55, yr, [CYAN, ORANGE, GREEN][i], desc, fs=18)
        if i < 2:
            arrow(ax, xs[i] + 3.7, 5.6, xs[i + 1] - 0.2, 5.6, MUTED)
    cats = [("NGS", 6.6, ORANGE), ("Personnel", 5.1, CYAN), ("CTC", 4.2, PURPLE), ("mIHC/GeoMx", 3.6, GREEN), ("Other", 10.5, GOLD)]
    total = sum(v for _, v, _ in cats)
    x0 = 1.0
    for name, val, color in cats:
        w = 13.5 * val / total
        ax.add_patch(Rectangle((x0, 2.0), w, 0.55, color=color, alpha=0.9))
        ax.text(x0 + w / 2, 1.45, name, color=TEXT, fontsize=11, ha="center")
        x0 += w
    ax.text(1.0, 2.85, "30억 / 3년 budget structure", color=TEXT, fontsize=15, fontweight="bold")
    ax.text(1.0, 0.75, "Preserve: serial blood + CTC phenotype + matched tissue-normal NGS.", color=CYAN, fontsize=16, fontweight="bold")
    save(fig, FIG / "slide11_budget_roadmap.png")


def slide12():
    fig, ax = setup_ax(FIG / "slide12_final_ask.png")
    ax.text(0.7, 7.5, "Fund a Science project that turns thyroidectomy into a liquid-biopsy perturbation experiment", color=TEXT, fontsize=22, fontweight="bold")
    center = (8.0, 4.7)
    ax.add_patch(Circle(center, 1.15, color=PURPLE, alpha=0.9))
    ax.text(*center, "Thyroidectomy", color=BG, fontsize=17, fontweight="bold", ha="center", va="center")
    nodes = [
        ("PTC", 3.0, 6.2, CYAN),
        ("CTC-EMT", 3.2, 3.0, ORANGE),
        ("matched NGS", 12.7, 6.2, GREEN),
        ("residual-risk\nbiology", 12.5, 3.0, GOLD),
    ]
    for lab, x, y, color in nodes:
        box(ax, x - 1.25, y - 0.45, 2.5, 0.9, lab, color, fs=14)
        arrow(ax, x, y, center[0] + (x - center[0]) * 0.18, center[1] + (y - center[1]) * 0.18, color)
    ax.text(1.2, 1.2, "갑상선절제술 전후 CTC-EMT 상태전이와 유전지도의 연속 해독을 통해,\n유두갑상선암의 잔존위험 생물학을 처음으로 시간축에서 해석한다.", color=TEXT, fontsize=18, fontweight="bold")
    save(fig, FIG / "slide12_final_ask.png")


SLIDES = [
    ("The Narrow Problem", "PTC is curable, but residual-risk biology is invisible after surgery", "Current postoperative tools cannot read the real-time residual tumor-cell state.", "slide01_problem.png"),
    ("Core Hypothesis", "Thyroidectomy as a human perturbation model", "Surgery should collapse tumor shedding; persistence or EMT-state shift after surgery is the signal.", "slide02_hypothesis.png"),
    ("Why Now", "CTC and EMT phenotypes are measurable in PTC", "The feasibility anchor exists, but the genetic map is missing.", "slide03_literature_gap.png"),
    ("The Missing Dataset", "No public dataset has the full serial CTC-EMT-NGS map", "This absence is the proposal's core justification.", "slide04_missing_dataset.png"),
    ("Public Pilot GO", "Driver mutation does not capture the full residual-risk state", "Local public pilot shows tissue states split beyond mutation group.", "slide05_driver_not_enough.png"),
    ("Tissue Context", "Spatial tissue states nominate where EMT/CTC signals may originate", "Spatial tissue states are organized, but not yet CTC proof.", "slide06_spatial_context.png"),
    ("Mechanistic Bridge", "From tissue state to circulating state", "Markers define CTC state and NGS interpretation.", "slide07_ctc_ngs_panel.png"),
    ("Prospective Design", "Serial blood around thyroidectomy", "The study converts surgery into a time-resolved perturbation experiment.", "slide08_sampling_timeline.png"),
    ("NGS Genetic Map", "Matched tissue NGS anchors the liquid-biopsy signal", "CTC phenotype alone can be nonspecific; NGS provides clone identity.", "slide09_genetic_map.png"),
    ("Endpoints", "What we will test, and what we will not claim", "Primary endpoint is biological state transition, not clinical deployment.", "slide10_claim_boundary.png"),
    ("Execution", "30억 / 3년 stage-gated execution", "Preserve serial blood, CTC phenotype, matched tissue-normal NGS.", "slide11_budget_roadmap.png"),
    ("Final Ask", "Turn thyroidectomy into a liquid-biopsy perturbation experiment", "The first PTC serial CTC-EMT genetic map after thyroidectomy.", "slide12_final_ask.png"),
]


def generate_figures() -> None:
    for fn in [slide01, slide02, slide03, slide04, slide05, slide06, slide07, slide08, slide09, slide10, slide11, slide12]:
        fn()


def add_textbox(slide, x, y, w, h, text, size=18, color=TEXT, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    run.font.name = "Arial"
    return tb


def add_rect(slide, x, y, w, h, fill=PANEL, line="#263244", transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill.replace("#", ""))
    shp.fill.transparency = transparency
    shp.line.color.rgb = RGBColor.from_string(line.replace("#", ""))
    shp.line.width = Pt(1.0)
    return shp


def add_rule(slide, x, y, w, color="#263244"):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.015))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color.replace("#", ""))
    shp.line.fill.background()
    return shp


def build_deck() -> Path:
    DECK.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    blank = prs.slide_layouts[6]
    for idx, (kicker, title, line, fig_name) in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(9, 13, 20)
        add_rule(slide, 0.0, 0.0, 16.0, CYAN)
        add_textbox(slide, 0.55, 0.28, 5.2, 0.35, "PTC CTC-EMT  |  SAMSUNG FUTURE TECHNOLOGY PROGRAM", 9, CYAN, True)
        add_textbox(slide, 12.15, 0.28, 3.25, 0.35, "SCIENCE TRACK  |  MECHANISM", 9, GOLD, True)
        add_textbox(slide, 0.6, 0.95, 0.95, 0.65, f"{idx:02d}", 28, TEXT, True)
        add_textbox(slide, 1.55, 0.78, 10.7, 0.55, title, 22, TEXT, True)
        add_textbox(slide, 1.57, 1.33, 10.8, 0.40, line, 12, MUTED, False)
        add_rect(slide, 0.6, 1.88, 14.85, 6.43, fill="#0d1420", line="#1f3348", transparency=0)
        slide.shapes.add_picture(str(FIG / fig_name), Inches(0.74), Inches(2.0), width=Inches(14.57), height=Inches(6.12))
        add_textbox(slide, 0.62, 8.45, 5.0, 0.24, "Proposal concept for Samsung Science Program", 8, MUTED, False)
        add_textbox(slide, 11.1, 8.45, 4.2, 0.24, "One disease | one perturbation | one transition | one genetic anchor", 8, MUTED, False)
    path = DECK / "kthyro_ctc_emt_samsung_science_pitch_v1.pptx"
    prs.save(path)
    prs.save(DECK / "kthyro_ctc_emt_samsung_science_pitch_v2_reference_style.pptx")
    return path


if __name__ == "__main__":
    generate_figures()
    path = build_deck()
    print(path)

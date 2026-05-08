"""Build Fig 12 — Thyroid hormone synthesis pathway as a beautiful PPTX."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from copy import deepcopy

OUT = Path(__file__).resolve().parent

# ---- color palette ----
RAI_FILL  = RGBColor(0xFF, 0xF0, 0xED); RAI_BD  = RGBColor(0xC0, 0x39, 0x2B)
NON_FILL  = RGBColor(0xEE, 0xF4, 0xFB); NON_BD  = RGBColor(0x2C, 0x5E, 0x9C)
TF_FILL   = RGBColor(0xEA, 0xF6, 0xEC); TF_BD   = RGBColor(0x3C, 0x6B, 0x4F)
COLLOID_F = RGBColor(0xFF, 0xF8, 0xD6); COLLOID_BD = RGBColor(0xB0, 0x8C, 0x1F)
BLOOD_F   = RGBColor(0xFF, 0xE1, 0xDE)
APICAL_F  = RGBColor(0xFC, 0xF3, 0xE7)
INK       = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT    = RGBColor(0x7B, 0x1F, 0x2A)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY      = RGBColor(0x66, 0x66, 0x66)

# ---- presentation: 16:9 widescreen, 13.333 x 7.5 inches ----
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

def add_text(slide, x, y, w, h, text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT,
             font="Noto Sans KR"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.margin_left=Inches(0); tf.margin_right=Inches(0)
    tf.margin_top=Inches(0.02); tf.margin_bottom=Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def add_box(slide, x, y, w, h, fill, line, line_w=1.5, rounded=True, dash=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    if dash:
        ln = shp.line._get_or_add_lnRef() if hasattr(shp.line, "_get_or_add_lnRef") else None
    shp.shadow.inherit = False
    # remove default text
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp

def gene_box(slide, x, y, w, h, name, sub, fill, line, name_size=14, sub_size=9):
    shp = add_box(slide, x, y, w, h, fill, line, line_w=2)
    tf = shp.text_frame; tf.word_wrap = True
    tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05)
    tf.margin_top=Inches(0.05); tf.margin_bottom=Inches(0.05)
    p1 = tf.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = name
    r1.font.name = "Noto Sans KR"; r1.font.size = Pt(name_size); r1.font.bold = True
    r1.font.color.rgb = line
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub
    r2.font.name = "Noto Sans KR"; r2.font.size = Pt(sub_size); r2.font.italic = True
    r2.font.color.rgb = GREY
    return shp

def add_arrow(slide, x1, y1, x2, y2, color=INK, weight=2.0, dashed=False):
    line = slide.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color; line.line.width = Pt(weight)
    # Add arrowhead via XML
    from pptx.oxml.ns import qn
    spPr = line.line._get_or_add_ln()
    tail = spPr.find(qn("a:tailEnd"))
    if tail is None:
        from lxml import etree
        tail = etree.SubElement(spPr, qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    if dashed:
        prstDash = spPr.find(qn("a:prstDash"))
        if prstDash is None:
            from lxml import etree
            prstDash = etree.SubElement(spPr, qn("a:prstDash"))
        prstDash.set("val", "dash")
    return line

# ===================== TITLE =====================
add_text(slide, 0.4, 0.12, 12.5, 0.4,
         "갑상선 호르몬 합성 + RAI 흡수 분자 pathway",
         size=24, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(slide, 0.4, 0.55, 12.5, 0.3,
         "8-gene RAI_8 (red) + THYROID_NONOVERLAP (blue) + TF backbone (green) 의 위치",
         size=12, color=GREY, align=PP_ALIGN.CENTER)

# ===================== Layer 1: TF backbone =====================
add_box(slide, 0.4, 1.0, 12.5, 0.95, TF_FILL, TF_BD, line_w=1.5)
add_text(slide, 0.6, 1.05, 8, 0.3,
         "REGULATORY LAYER · Lineage transcription factors (TF backbone)",
         size=11, bold=True, color=TF_BD)
gene_box(slide, 0.7, 1.40, 1.6, 0.5, "FOXE1", "TTF-2 / FKHL15", WHITE, TF_BD, name_size=12, sub_size=8.5)
gene_box(slide, 2.5, 1.40, 1.6, 0.5, "NKX2-1", "TTF-1", WHITE, TF_BD, name_size=12, sub_size=8.5)
gene_box(slide, 4.3, 1.40, 1.6, 0.5, "PAX8", "PAX-related", WHITE, TF_BD, name_size=12, sub_size=8.5)
gene_box(slide, 6.1, 1.40, 1.6, 0.5, "HHEX", "endoderm TF", WHITE, TF_BD, name_size=12, sub_size=8.5)
add_text(slide, 8.0, 1.50, 4.7, 0.3,
         "→ NIS / TG / TPO / TSHR promoter 직접 binding",
         size=10, color=TF_BD, font="Noto Sans KR")

# downward driving arrow
add_arrow(slide, 6.7, 1.95, 6.7, 2.45, TF_BD, 1.5)
add_text(slide, 6.85, 2.05, 1.5, 0.25, "drives", size=10, color=TF_BD)

# ===================== Layer 2: Follicle (middle) =====================
# Blood (basolateral) band
add_box(slide, 0.4, 3.05, 4.4, 0.30, BLOOD_F, RAI_BD, line_w=1)
add_text(slide, 0.6, 3.07, 4.0, 0.25, "BLOOD  (basolateral)", size=10, bold=True, color=RAI_BD)
# Apical band
add_box(slide, 8.5, 3.05, 4.4, 0.30, APICAL_F, COLLOID_BD, line_w=1)
add_text(slide, 8.7, 3.07, 4.0, 0.25, "APICAL membrane", size=10, bold=True, color=COLLOID_BD)

# Colloid lumen — circle
shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.7), Inches(2.7), Inches(2.0), Inches(2.0))
shp.fill.solid(); shp.fill.fore_color.rgb = COLLOID_F
shp.line.color.rgb = COLLOID_BD; shp.line.width = Pt(2.5)
shp.text_frame.text = ""
add_text(slide, 5.8, 2.95, 1.8, 0.3, "Colloid", size=14, bold=True,
         color=COLLOID_BD, align=PP_ALIGN.CENTER)
add_text(slide, 5.8, 3.25, 1.8, 0.25, "(lumen)", size=10, color=COLLOID_BD, align=PP_ALIGN.CENTER)
add_text(slide, 5.8, 3.65, 1.8, 0.3, "TG-I  (iodinated)", size=12, bold=True,
         color=RAI_BD, align=PP_ALIGN.CENTER)

# TSHR (RAI_8) — left side basolateral
gene_box(slide, 0.6, 3.55, 1.4, 0.6, "TSHR", "(TSH 수용체)", RAI_FILL, RAI_BD)

# I- arrow into NIS
add_text(slide, 1.2, 3.45, 0.5, 0.3, "I⁻", size=18, bold=True, color=RAI_BD)
gene_box(slide, 3.6, 3.55, 1.4, 0.6, "NIS", "(SLC5A5)", RAI_FILL, RAI_BD)
add_arrow(slide, 2.1, 3.4, 3.6, 3.85, RAI_BD, 2.5)

# I- arrow from cell into colloid via Pendrin
gene_box(slide, 8.3, 3.55, 1.5, 0.6, "Pendrin", "(SLC26A4)", NON_FILL, NON_BD)
# arrow from cell to Pendrin to colloid
add_arrow(slide, 7.9, 3.85, 8.3, 3.85, NON_BD, 2.0)
add_arrow(slide, 9.8, 3.85, 7.7, 3.85, NON_BD, 2.0)

# DIO1 (RAI_8) and DIO2 (NONOVERLAP)
gene_box(slide, 11.0, 2.95, 1.7, 0.6, "DIO1", "(deiodinase 1)", RAI_FILL, RAI_BD)
gene_box(slide, 11.0, 3.65, 1.7, 0.6, "DIO2", "(deiodinase 2)", NON_FILL, NON_BD)

# T4/T3 output
add_arrow(slide, 11.0, 3.25, 13.1, 2.5, COLLOID_BD, 3.0)
add_text(slide, 11.4, 1.95, 2.0, 0.4, "T4 / T3 분비 →", size=14, bold=True, color=COLLOID_BD)
add_text(slide, 11.4, 2.30, 2.0, 0.3, "(혈류로)", size=11, color=COLLOID_BD)

# ===================== Layer 3: Effector enzymes (bottom) =====================
gene_box(slide, 0.6, 5.0, 1.6, 0.7, "TFF3", "(secretory)", NON_FILL, NON_BD, name_size=13)
gene_box(slide, 2.4, 5.0, 1.6, 0.7, "GLIS3", "(supportive TF)", NON_FILL, NON_BD, name_size=13)
gene_box(slide, 4.4, 5.0, 1.6, 0.7, "TG", "(thyroglobulin)", RAI_FILL, RAI_BD, name_size=13)
gene_box(slide, 6.2, 5.0, 1.6, 0.7, "TPO", "(peroxidase)", RAI_FILL, RAI_BD, name_size=13)
gene_box(slide, 8.0, 5.0, 1.8, 0.7, "DUOX1 / 2", "(H₂O₂ supply)", NON_FILL, NON_BD, name_size=13)
gene_box(slide, 10.0, 5.0, 1.6, 0.7, "IYD", "(I⁻ recycle)", NON_FILL, NON_BD, name_size=13)

# arrows: DUOX → TPO (H2O2), TPO → colloid (oxidation), TG → colloid (substrate)
add_arrow(slide, 8.0, 5.3, 7.8, 5.3, NON_BD, 1.5)
add_text(slide, 7.0, 4.8, 1.0, 0.3, "H₂O₂", size=9, color=NON_BD)
add_arrow(slide, 7.0, 5.0, 6.7, 4.7, RAI_BD, 1.5)
add_text(slide, 6.6, 4.65, 1.6, 0.25, "I⁻ 산화", size=9, color=RAI_BD)
add_arrow(slide, 5.2, 5.0, 6.0, 4.7, RAI_BD, 1.5, dashed=True)

# ===================== Legend at bottom =====================
add_box(slide, 0.4, 6.4, 12.5, 0.85, RGBColor(0xFA,0xFA,0xFA), GREY, line_w=0.5)
def legend_chip(x, y, fill, line, label):
    add_box(slide, x, y, 0.5, 0.32, fill, line, line_w=2)
    add_text(slide, x+0.6, y+0.05, 5.5, 0.25, label, size=11, bold=True, color=line)
legend_chip(0.7, 6.55, RAI_FILL, RAI_BD, "RAI_8 panel — 8 genes (main deployable readout)")
legend_chip(0.7, 6.95, NON_FILL, NON_BD, "THYROID_NONOVERLAP — 8 genes (zero overlap with RAI_8)")
legend_chip(8.0, 6.55, TF_FILL, TF_BD,  "TF backbone — Damante 4-TF lineage")
legend_chip(8.0, 6.95, COLLOID_F, COLLOID_BD, "Hormone output (T4 / T3)")

# ===================== Footer note =====================
add_text(slide, 0.4, 7.30, 12.5, 0.18,
         "★ DM1 = 위 모든 layer (TF + RAI_8 + NONOVERLAP) 가 동시에 침묵 → RAI 흡수 machinery 시스템적 lineage collapse",
         size=10, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# ---- save ----
out_pptx = OUT/"fig12_thyroid_pathway.pptx"
prs.save(str(out_pptx))
print(f"[done] {out_pptx}  ({out_pptx.stat().st_size:,} B)")
